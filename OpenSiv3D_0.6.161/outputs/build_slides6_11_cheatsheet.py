from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = r"C:\Users\tomat\source\repos\RyotaIzumi\I-wanna-Siv3D\OpenSiv3D_0.6.161\outputs"
OUT = ROOT + r"\slides6_11_english_cheatsheet_one_page.docx"

sections = [
    ("6  Constraint workflow", "We use four steps: describe possible safe routes; encode the design rules as 3SAT or a graph; solve for a satisfying assignment or Hamilton path; then convert the solution into timed stage elements. The solver does not replace creative design. It acts as a scaffold that protects at least one valid route while we still control the music and presentation."),
    ("7  3SAT prototype", "The player assigns each variable by entering TRUE or FALSE before the timer ends. The game then checks each clause. A clause is passable when at least one of its three literals is true. Boolean assignments therefore become movement decisions, while clause satisfaction becomes a playable obstacle check."),
    ("8  Why the mapping matters", "The goal is formal correspondence, not only visual similarity. Variables become forced choices, literals become route-linked conditions, and clauses become pass-or-fail gates. If the mapping is correct, the stage is clearable exactly when the represented formula is satisfiable. Proving and testing this relationship is a central research task."),
    ("9  Live 3SAT demonstration", "Switch to the game. Show one timed variable assignment, then show how the clause checks respond. Intentionally demonstrate a failing assignment once; replay with a satisfying assignment and clear the sequence. Ask the audience to notice whether the relationship between the logical choice and the game response is understandable. Then return to the slides."),
    ("10  Hamilton Path prototype", "Each yellow platform represents a vertex, and permitted transitions represent graph edges. The player must visit every platform exactly once before reaching the green goal. Repeated visits and invalid transitions must be blocked or punished. A Hamilton path therefore provides a complete safe route. Integrating this prototype into a real barrage is the next implementation task."),
    ("11  Closing and next steps", "Next, implement a practical Hamilton Path example; improve the 3SAT example's readability, pacing, difficulty, and music synchronization; and evaluate both logical correctness and playability. The central question is: can formal constraints help us create expressive bullet hells without losing a guaranteed route? Invite feedback on both the formal model and the game design."),
]

def set_font(run, size, bold=False, color="111111"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.48)
sec.bottom_margin = Inches(0.45)
sec.left_margin = Inches(0.55)
sec.right_margin = Inches(0.55)
sec.header_distance = Inches(0.25)
sec.footer_distance = Inches(0.25)

# Two-column A4 is a named compact-print override.
sectPr = sec._sectPr
cols = sectPr.xpath("./w:cols")
cols_el = cols[0] if cols else OxmlElement("w:cols")
if not cols:
    sectPr.append(cols_el)
cols_el.set(qn("w:num"), "2")
cols_el.set(qn("w:space"), "420")
cols_el.set(qn("w:sep"), "1")

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.08

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Slides 6–11 | English Presenter Cheat Sheet")
set_font(r, 16, True, "1F4D78")
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(7)
r = p2.add_run("3SAT and Hamilton Path for safer bullet-hell generation · one-page print edition")
set_font(r, 8.8, False, "555A61")

for title, body in sections:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1.5)
    r = p.add_run(title)
    set_font(r, 11.2, True, "2E74B5")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(body)
    set_font(r, 9.5, False, "111111")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Speaking cue: pause after each key claim; emphasize the words in the slide title.")
set_font(r, 7.5, False, "777777")

doc.core_properties.title = "Slides 6–11 English Presenter Cheat Sheet"
doc.core_properties.subject = "One-page speaking notes"
doc.save(OUT)
print(OUT)
